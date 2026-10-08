"""
Validación de incidencias del cliente (v24): los cálculos.

La pantalla vive en components/cliente.py y la lectura en data.py. Aquí solo
hay pandas, igual que en director_calc.py: cada función recibe tablas ya leídas
y regresa diccionarios listos para pintar, así que todo se prueba contra los
CSV de cualquier corrida sin levantar la app.

Lo que el cliente decide NO se aplica aquí. El cierre (v7.3) sigue leyendo el
Excel de INCIDENCIAS FINALES, y este módulo lo arma con el mismo formato con
que se cerró agosto.

La pantalla le cuenta al cliente, en palabras, las reglas con que el cierre
aplica una incidencia aprobada (celda v7.3 del cuaderno):
  - Da por cumplido SOLO el KPI que reclamó y SOLO en esa semana.
  - SOS: la semana sube al objetivo de la tienda; nunca baja un valor que ya
    iba mejor. Si la tienda no trae objetivo en el AOP, la categoría se da por
    cumplida en el mes.
  - Exhibiciones: a la semana se le completan los puntos que le faltaban para
    su objetivo (4, o 2 en Sams y City Club).
  - OOS: las encuestas sin contestar de esa tienda y semana se dan por
    contestadas.
  - Una semana con incidencia de visita de TRAX no se toca: no entra al
    promedio de todos modos.
Si esas reglas cambian en el cuaderno, hay que cambiar hechos() también.
"""
import io
import json
import re
from datetime import timedelta, timezone

import pandas as pd

from components import cumplimiento

PENDIENTE, APROBADA, NO_APROBADA = 'PENDIENTE', 'APROBADA', 'NO_APROBADA'
ESTADOS_CLIENTE = (PENDIENTE, APROBADA, NO_APROBADA)

KPIS = ('SOS WHISKY', 'SOS TEQUILA', 'SOS VODKA', 'EXHIBICIONES', 'OOS')
KPI_NOMBRE = {'SOS WHISKY': 'SOS whisky', 'SOS TEQUILA': 'SOS tequila', 'SOS VODKA': 'SOS vodka',
              'EXHIBICIONES': 'Exhibiciones', 'OOS': 'OOS'}
KPI_CORTO = {'SOS WHISKY': 'whisky', 'SOS TEQUILA': 'tequila', 'SOS VODKA': 'vodka',
             'EXHIBICIONES': 'exhibiciones', 'OOS': 'OOS'}
# Por KPI: columna del valor, columna del objetivo (las dos en la escala de la
# base: SOS en 0-100, exhibiciones en puntos) y su llave en cumplimiento.categorias().
COLUMNAS_KPI = {
    'SOS WHISKY': ('sos_whisky', 'obj_whisky', 'Whisky'),
    'SOS TEQUILA': ('sos_tequila', 'obj_tequila', 'Tequila'),
    'SOS VODKA': ('sos_vodka', 'obj_vodka', 'Vodka'),
    'EXHIBICIONES': ('exh_puntos', 'obj_exh', 'EXH'),
}

# Por qué no se aprueba. Lista corta a propósito: se puede contar en el resumen.
# "Otro" obliga a escribir el comentario.
MOTIVOS_NO = ('La foto no muestra el problema', 'Esa semana ya cumplía', 'TRAX sí lo reconoce',
              'Evidencia borrosa o incompleta', 'No corresponde al KPI', 'Duplicada', 'Otro')
MOTIVO_OTRO = 'Otro'

SEMANAS_TIRA = 12        # la tira de seguimiento enseña a lo más 12 semanas
EMPATE = cumplimiento.EMPATE
MESES_CORTOS = ('ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep', 'oct', 'nov', 'dic')
# Hora de la Ciudad de México. Sin horario de verano desde 2022; se fija el
# desfase para no depender de la base de zonas horarias del servidor.
HORA_MEXICO = timezone(timedelta(hours=-6))

# El Excel para el cierre: las columnas de la tabla incidencias en el orden del
# archivo de agosto, VALIDACION FINAL y, al final, lo nuevo del cliente. El
# cuaderno solo lee curt, semana, tipo, periodo_id y VALIDACION FINAL.
COLUMNAS_EXCEL = ('id', 'curt', 'ruta', 'periodo_id', 'tienda', 'cadena', 'canal', 'tipo', 'semana',
                  'comentario', 'fotos', 'reportada_por', 'created_at', 'link_trax', 'estado',
                  'resuelta_por', 'resuelta_en', 'motivo_rechazo', 'incidencia', 'categoria', 'productos')
COLUMNA_VALIDACION = 'VALIDACION FINAL'
COLUMNAS_EXCEL_CLIENTE = ('motivo_cliente', 'comentario_cliente', 'validada_por', 'validada_en')
# Los mismos textos que trae la columna en el archivo de agosto.
TEXTO_EXCEL = {APROBADA: 'Aprobada', NO_APROBADA: 'No aprobada', PENDIENTE: None}

_TEXTOS = ('ruta', 'periodo_id', 'tienda', 'cadena', 'canal', 'comentario', 'reportada_por', 'created_at',
           'link_trax', 'resuelta_por', 'resuelta_en', 'motivo_rechazo', 'incidencia', 'categoria',
           'validada_por', 'validada_en', 'motivo_cliente', 'comentario_cliente')
_URL = re.compile(r'https?://[^\s<>"\']+', re.IGNORECASE)


# ------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------
def texto(v) -> str:
    """El texto limpio, o '' si no hay. NaN es truthy: (v or '') no lo atrapa."""
    if v is None or isinstance(v, (list, tuple, dict)):
        return ''
    try:
        if pd.isna(v):
            return ''
    except (TypeError, ValueError):
        return ''
    return str(v).strip()


def numero(v):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return None if pd.isna(v) else v


def si(v) -> bool:
    return str(v).strip().lower() in ('1', '1.0', 'true')


def curt_txt(v) -> str:
    """La CURT como en la base (TEXT): '2744', no '2744.0'."""
    n = numero(v)
    if n is not None and float(n).is_integer():
        return str(int(n))
    return texto(v)


def lista(v) -> list:
    """fotos y productos: llegan como lista (Supabase), texto JSON (Excel o CSV)
    o arreglo de Postgres ({a,b})."""
    if v is None:
        return []
    if hasattr(v, 'tolist') and not isinstance(v, str):
        v = v.tolist()
    if isinstance(v, (list, tuple)):
        return [texto(x) for x in v if texto(x)]
    t = texto(v)
    if not t:
        return []
    if t.startswith('['):
        try:
            return lista(json.loads(t))
        except ValueError:
            pass
    if t.startswith('{') and t.endswith('}'):
        return [x.strip().strip('"') for x in t[1:-1].split(',') if x.strip().strip('"')]
    return [t]


def tipo_kpi(v) -> str:
    """'EXH' se renombró a 'EXHIBICIONES' en v13; quedan incidencias viejas."""
    t = texto(v).upper()
    return 'EXHIBICIONES' if t == 'EXH' else t


def folio(inc_id) -> str:
    return f"INC-{int(inc_id):04d}"


def url_trax(v) -> str:
    """El link de la visita. El promotor a veces pega el mensaje completo de
    TRAX ("Incidencia de Walmart ... https://..."): se toma solo la URL."""
    m = _URL.search(texto(v))
    return m.group(0).rstrip('.,;)]') if m else ''


def fecha_corta(v) -> str:
    """'12 sep · 18:40', en hora de la Ciudad de México."""
    t = pd.to_datetime(v, utc=True, errors='coerce')
    if pd.isna(t):
        return ''
    t = t.tz_convert(HORA_MEXICO)
    return f"{t.day} {MESES_CORTOS[t.month - 1]} · {t:%H:%M}"


def valor_txt(tipo, v) -> str:
    if v is None:
        return '—'
    if tipo == 'EXHIBICIONES':
        return f"{v:g} pts" if float(v).is_integer() else f"{v:.2f} pts"
    return f"{v:.2f}%"


def puntos(v) -> str:
    return '1 punto' if v == 1 else f"{v:g} puntos"


def objetivo_txt(tipo, v) -> str:
    """El objetivo como lo guarda la base: 22.4%, 1.18%, 4 pts."""
    if v is None:
        return '—'
    return f"{v:g} pts" if tipo == 'EXHIBICIONES' else f"{v:g}%"


def mes_corto(periodos: pd.DataFrame, periodo_id) -> str:
    fila = periodos[periodos['periodo_id'] == periodo_id]
    if len(fila) == 0:
        return str(periodo_id)
    return str(fila.iloc[0].get('mes') or periodo_id).strip().capitalize()[:3]


# ------------------------------------------------------------
# Incidencias
# ------------------------------------------------------------
def preparar(inc: pd.DataFrame) -> pd.DataFrame:
    """Incidencias listas para la pantalla: tipos parejos y sin NaN en los textos.

    Sirve igual para el periodo completo (todas las columnas) que para la
    lectura ligera de todos los meses: la columna que no venga queda vacía.
    """
    df = inc.copy() if inc is not None and len(inc) else pd.DataFrame(columns=['id'])
    for c in ('id', 'curt', 'tipo', 'semana', 'estado', 'validacion_cliente', 'fotos', 'productos') + _TEXTOS:
        if c not in df.columns:
            df[c] = None
    df['id'] = pd.to_numeric(df['id'], errors='coerce')
    df = df.dropna(subset=['id']).copy()
    df['id'] = df['id'].astype(int)
    df['curt'] = df['curt'].map(curt_txt)
    df['tipo'] = df['tipo'].map(tipo_kpi)
    df['semana'] = pd.to_numeric(df['semana'], errors='coerce').astype('Int64')
    for c in _TEXTOS:
        df[c] = df[c].map(texto)
    df['estado'] = df['estado'].map(lambda v: texto(v).upper() or 'PENDIENTE')
    df['validacion_cliente'] = df['validacion_cliente'].map(
        lambda v: texto(v).upper() if texto(v).upper() in ESTADOS_CLIENTE else PENDIENTE)
    df['fotos'] = df['fotos'].map(lista)
    df['productos'] = df['productos'].map(lista)
    df['creada'] = pd.to_datetime(df['created_at'], utc=True, errors='coerce')
    df['folio'] = df['id'].map(folio)
    df['url_trax'] = df['link_trax'].map(url_trax)
    return df.reset_index(drop=True)


RANGO_SUPERVISOR = {'AUTORIZADA': 0, 'PENDIENTE': 1, 'NO_AUTORIZADA': 2}


def ordenar(df: pd.DataFrame) -> pd.DataFrame:
    """El orden de la cola: primero lo que el supervisor autorizó, luego lo que
    no ha revisado y al final lo que rechazó. Dentro, por tienda, KPI y semana,
    para que la serie de una misma tienda quede junta."""
    if len(df) == 0:
        return df
    llave = df['estado'].map(RANGO_SUPERVISOR).fillna(1)
    return (df.assign(_r=llave, _t=df['tienda'].str.upper())
              .sort_values(['_r', '_t', 'curt', 'tipo', 'semana', 'id'])
              .drop(columns=['_r', '_t']).reset_index(drop=True))


def orden_periodos(periodos: pd.DataFrame) -> list:
    """Los periodo_id del más viejo al más nuevo."""
    if len(periodos) == 0:
        return []
    p = periodos.copy()
    p['_f'] = pd.to_datetime(p.get('fecha_inicio'), errors='coerce')
    return p.sort_values(['_f', 'periodo_id'])['periodo_id'].tolist()


def semanas_en_orden(periodos: pd.DataFrame) -> list:
    """Todas las (periodo, semana) de todos los periodos, en orden.

    Se ordena por periodo y no solo por número de semana: el calendario de TRAX
    vuelve a empezar en enero.
    """
    salida = []
    por_id = periodos.set_index('periodo_id') if len(periodos) else periodos
    for pid in orden_periodos(periodos):
        fila = por_id.loc[pid]
        ini, fin = numero(fila.get('semana_inicio')), numero(fila.get('semana_fin'))
        if ini is None or fin is None or fin < ini:
            continue
        salida += [(pid, s) for s in range(int(ini), int(fin) + 1)]
    return salida


def marcas_serie(todas: pd.DataFrame) -> pd.DataFrame:
    """Por incidencia (índice = id): en cuántas semanas distintas se ha reclamado
    la misma tienda y el mismo KPI, contando meses anteriores ('semanas_serie'),
    y cuántas incidencias hay de esa misma tienda, KPI y semana ('iguales')."""
    if len(todas) == 0:
        return pd.DataFrame(columns=['semanas_serie', 'iguales'])
    t = todas.dropna(subset=['semana'])
    semanas = (t.drop_duplicates(['curt', 'tipo', 'periodo_id', 'semana'])
                .groupby(['curt', 'tipo']).size().rename('semanas_serie'))
    iguales = t.groupby(['curt', 'tipo', 'periodo_id', 'semana']).size().rename('iguales')
    out = (todas[['id', 'curt', 'tipo', 'periodo_id', 'semana']]
           .join(semanas, on=['curt', 'tipo']).join(iguales, on=['curt', 'tipo', 'periodo_id', 'semana']))
    out[['semanas_serie', 'iguales']] = out[['semanas_serie', 'iguales']].fillna(1).astype(int)
    return out.set_index('id')[['semanas_serie', 'iguales']]


def serie_de(todas: pd.DataFrame, curt, tipo) -> pd.DataFrame:
    """Todas las incidencias de la misma tienda y KPI, de todos los meses."""
    s = todas[(todas['curt'] == curt_txt(curt)) & (todas['tipo'] == tipo)]
    return s.sort_values(['creada', 'id'])


# ------------------------------------------------------------
# Datos del KPI
# ------------------------------------------------------------
def valores_semana(ds: pd.DataFrame, os_: pd.DataFrame, tipo) -> dict:
    """{(periodo, semana): lo que midió TRAX ese KPI esa semana}.

    Colores como los cuadritos del detalle de tienda: SOS contra su objetivo;
    exhibiciones con el CUMPLIO 4 de la semana, que ya trae lo que completan
    licor y ron. 'llega' es otra cosa: si los puntos o el SOS alcanzan SOLOS el
    objetivo, que es lo único que una incidencia aprobada puede subir.
    """
    out = {}
    if tipo == 'OOS':
        for _, w in (os_ if os_ is not None else pd.DataFrame()).iterrows():
            if numero(w.get('semana')) is None:
                continue
            obj = int(numero(w.get('obj_oos')) or 0)
            nc = int(numero(w.get('no_cont_oos')) or 0)
            cont = int(numero(w.get('contestadas_oos')) or 0)
            out[(w.get('periodo_id'), int(w.get('semana')))] = {
                'obj': obj, 'cont': cont, 'nc': nc, 'sin_obj': obj <= 0,
                'ok': None if obj <= 0 else nc == 0}
        return out
    col, col_obj, _ = COLUMNAS_KPI[tipo]
    for _, w in (ds if ds is not None else pd.DataFrame()).iterrows():
        if numero(w.get('semana')) is None:
            continue
        clave = (w.get('periodo_id'), int(w.get('semana')))
        v, obj = numero(w.get(col)), numero(w.get(col_obj))
        sin_obj = obj is None or obj <= 0
        if si(w.get('incidencia')):
            out[clave] = {'trax': True, 'obj': obj, 'sin_obj': sin_obj, 'ok': None}
            continue
        if v is None:
            out[clave] = {'sin_dato': True, 'obj': obj, 'sin_obj': sin_obj, 'ok': None}
            continue
        if sin_obj:
            ok, llega = None, False
        elif tipo == 'EXHIBICIONES':
            ok = si(w.get('cumplio_4'))
            llega = v >= obj - 1e-9
        else:
            ok = v >= obj - 1e-9
            llega = v >= obj + EMPATE
        out[clave] = {'v': v, 'obj': obj, 'ok': ok, 'llega': llega, 'sin_obj': sin_obj,
                      'empata': (not sin_obj and tipo != 'EXHIBICIONES' and abs(v - obj) < EMPATE)}
    return out


def para_cumplimiento(rt) -> dict:
    """Una fila de resumen_tienda con los nombres y escalas de adaptar_tiendas()
    (SOS en 0-1, objetivos en 0-100), que es lo que espera cumplimiento."""
    g = rt.get

    def sos(c):
        v = numero(g(c))
        return None if v is None else v / 100.0

    return {'Total Whisky': sos('sos_whisky'), 'Total tequila': sos('sos_tequila'),
            'Total vodka': sos('sos_vodka'), 'Objetivo Whisky': g('obj_whisky'),
            'Objetivo Tequila': g('obj_tequila'), 'Objetivo Vodka': g('obj_vodka'),
            'cumplio_wtv': g('cumplio_wtv'), 'CUMPLIO 4': g('cumplio_4'),
            'Puntos Promedio Exhibición': g('exh_puntos'), 'Objetivo Puntos HS': g('obj_exh'),
            'Cadena': g('cadena')}


def mes_kpi(rt, os_mes: pd.DataFrame, tipo) -> dict:
    """El KPI del mes como lo dejó la última corrida cargada, y si cumple.

    El veredicto es el del pipeline (components/cumplimiento.py), el mismo que
    ven el promotor y el supervisor en su app: aquí no se vuelve a restar nada.
    """
    if tipo == 'OOS':
        o = os_mes if os_mes is not None else pd.DataFrame()
        obj = int(pd.to_numeric(o.get('obj_oos'), errors='coerce').fillna(0).sum()) if len(o) else 0
        nc = int(pd.to_numeric(o.get('no_cont_oos'), errors='coerce').fillna(0).sum()) if len(o) else 0
        cont = int(pd.to_numeric(o.get('contestadas_oos'), errors='coerce').fillna(0).sum()) if len(o) else 0
        return {'obj': obj, 'cont': cont, 'nc': nc, 'sin_obj': obj <= 0,
                'ok': None if obj <= 0 else nc == 0, 'nota': ''}
    if rt is None:
        return None
    col, col_obj, cat = COLUMNAS_KPI[tipo]
    fila = para_cumplimiento(rt)
    cumple = cumplimiento.categorias(fila)[cat]
    v, obj = numero(rt.get(col)), numero(rt.get(col_obj))
    sin_obj = obj is None or obj <= 0
    nota = ''
    if tipo == 'EXHIBICIONES':
        nota = cumplimiento.nota_exh(fila)
    elif cumple and sin_obj:
        nota = 'Cumple por una incidencia aprobada: la tienda no trae objetivo en el AOP.'
    elif sin_obj:
        nota = 'Sin objetivo en el AOP: así la categoría no se puede cumplir.'
    elif not cumple and v is not None and v >= obj - 1e-9:
        nota = 'Empata con el objetivo a 2 decimales, pero con todos sus decimales queda abajo.'
    return {'v': v, 'obj': obj, 'ok': cumple, 'sin_obj': sin_obj, 'nota': nota}


def historia_meses(rt_todos: pd.DataFrame, os_todos: pd.DataFrame, tipo, periodos: pd.DataFrame,
                   hasta, maximo=4) -> list:
    """El KPI mes por mes, del más viejo al más nuevo, hasta el periodo 'hasta'."""
    orden = orden_periodos(periodos)
    if hasta not in orden:
        return []
    fin = orden.index(hasta)
    salida = []
    for pid in orden[max(0, fin - maximo + 1): fin + 1]:
        if tipo == 'OOS':
            o = os_todos[os_todos['periodo_id'] == pid] if len(os_todos) else os_todos
            if len(o) == 0:
                continue
            m = mes_kpi(None, o, tipo)
            if m['sin_obj']:
                continue
            txt = 'todas contestadas' if m['nc'] == 0 else f"{m['nc']} sin contestar de {m['obj']}"
        else:
            f = rt_todos[rt_todos['periodo_id'] == pid] if len(rt_todos) else rt_todos
            if len(f) == 0:
                continue
            m = mes_kpi(f.iloc[0], None, tipo)
            if m['v'] is None:
                continue
            txt = f"{valor_txt(tipo, m['v'])} de {objetivo_txt(tipo, m['obj'])}" if not m['sin_obj'] \
                else f"{valor_txt(tipo, m['v'])}, sin objetivo"
        salida.append({'periodo': pid, 'mes': mes_corto(periodos, pid), 'texto': txt, 'ok': m['ok'],
                       'actual': pid == hasta})
    return salida


def hechos(tipo, semana, sem: dict, mes: dict, es_ps: bool) -> list:
    """Lo que conviene saber antes de decidir, en palabras: [(tono, texto)].

    Solo hechos de los datos y las reglas del cierre, nunca una estimación.
    tono: 'nada' = aprobarla no cambia nada; 'info' = contexto.
    Se usa con las incidencias por validar: con un mes ya cerrado, la corrida
    cargada ya trae lo aprobado y estos avisos dejarían de ser ciertos.
    """
    cat = KPI_CORTO[tipo]
    s = f"la semana {semana}"
    if sem is None:
        return [('info', f"No hay datos de {cat} de esta tienda en {s}.")]
    if sem.get('trax'):
        return [('nada', f"En {s} TRAX marcó incidencia de visita: esa semana no entra al promedio "
                         f"y el cierre no la toca.")]
    if sem.get('sin_dato'):
        return [('info', f"No hay datos de {cat} de esta tienda en {s}.")]
    salida = []
    if tipo == 'OOS':
        if sem['sin_obj']:
            salida.append(('nada', f"En {s} la tienda no tenía encuestas de agotados."))
        elif sem['nc'] == 0:
            salida.append(('nada', f"En {s} no quedó ninguna encuesta sin contestar: aprobarla no cambia nada."))
        else:
            salida.append(('info', f"En {s} quedaron {sem['nc']} encuestas sin contestar de {sem['obj']}. "
                                   f"Si se aprueba, se dan por contestadas."))
        return salida
    if sem['sin_obj']:
        if tipo == 'EXHIBICIONES':
            salida.append(('info', "La tienda no tiene objetivo de exhibiciones."))
        else:
            salida.append(('info', f"La tienda no tiene objetivo de {cat} en el AOP. Si se aprueba, "
                                   f"el cierre da el {cat} por cumplido todo el mes."))
        return salida
    v, obj = sem['v'], sem['obj']
    if sem['llega']:
        if tipo == 'EXHIBICIONES':
            salida.append(('nada', f"En {s} ya tenía sus {puntos(obj)} ({v:g}): aprobarla no cambia nada."))
        else:
            salida.append(('nada', f"En {s} ya cumplía: {valor_txt(tipo, v)} contra {objetivo_txt(tipo, obj)}. "
                                   f"Aprobarla no cambia nada."))
    elif tipo == 'EXHIBICIONES':
        salida.append(('info', f"En {s} tuvo {puntos(v)} de {obj:g}. Si se aprueba, se le completan los que faltan."))
    elif sem.get('empata'):
        salida.append(('info', f"En {s} empata con su objetivo a 2 decimales ({valor_txt(tipo, v)}). "
                               f"Si se aprueba, cuenta como cumplida."))
    else:
        salida.append(('info', f"En {s} quedó en {valor_txt(tipo, v)}, abajo de su objetivo de "
                               f"{objetivo_txt(tipo, obj)}. Si se aprueba, esa semana cuenta como cumplida."))
    if mes and mes.get('ok') and not mes.get('sin_obj'):
        quien = 'Las exhibiciones del mes ya cumplen' if tipo == 'EXHIBICIONES' else f"El {cat} del mes ya cumple"
        salida.append(('nada', f"{quien}: aprobarla no cambia el resultado de la tienda."))
    elif es_ps:
        salida.append(('nada', "La tienda ya es Perfect Store este mes."))
    return salida


def tira(serie: pd.DataFrame, valores: dict, semanas: list, id_actual, maximo=SEMANAS_TIRA) -> dict:
    """La tira semana por semana de una serie (misma tienda y KPI).

    Va de la primera semana con incidencia a la última que tenga datos o
    incidencia, así se ve también si después mejoró. Si no cabe, se queda con
    las últimas 'maximo' y cuenta cuántas semanas con incidencia quedaron fuera.
    """
    pos = {k: i for i, k in enumerate(semanas)}
    claves = {(p, int(s)) for p, s in zip(serie['periodo_id'], serie['semana']) if pd.notna(s)}
    con_inc = sorted(pos[k] for k in claves if k in pos)
    if not con_inc:
        return {'celdas': [], 'fuera': 0}
    fin = max([con_inc[-1]] + [pos[k] for k in valores if k in pos])
    ini = con_inc[0]
    fuera = 0
    if fin - ini + 1 > maximo:
        ini = fin - maximo + 1
        fuera = sum(1 for i in con_inc if i < ini)
    celdas = []
    for i in range(ini, fin + 1):
        p, s = semanas[i]
        incs = serie[(serie['periodo_id'] == p) & (serie['semana'] == s)]
        celdas.append({'periodo': p, 'semana': s, 'valor': valores.get((p, s)),
                       'incidencias': incs.to_dict('records'), 'actual': bool((incs['id'] == id_actual).any())})
    return {'celdas': celdas, 'fuera': fuera}


def lecturas(serie: pd.DataFrame, t: dict, tipo, motivo_actual='') -> list:
    """La serie en palabras."""
    n_sem = serie.drop_duplicates(['periodo_id', 'semana']).shape[0]
    cuenta = serie['validacion_cliente'].value_counts()
    ap, no, pe = (int(cuenta.get(k, 0)) for k in (APROBADA, NO_APROBADA, PENDIENTE))
    partes = [f"{n} {txt}" for n, txt in ((ap, 'aprobada' if ap == 1 else 'aprobadas'),
                                          (no, 'no aprobada' if no == 1 else 'no aprobadas'),
                                          (pe, 'por validar')) if n]
    semanas = [c for c in t['celdas'] if c['incidencias']]
    rango = f"la S{semanas[0]['semana']} y la S{semanas[-1]['semana']}" if semanas else ''
    salida = [f"Se reclamó {KPI_NOMBRE[tipo]} en {n_sem} semanas distintas"
              + (f" entre {rango}" if rango else '') + ': ' + _y(partes) + '.']
    if t['fuera']:
        salida.append(f"Hay {t['fuera']} semana(s) con incidencia antes de las que se ven aquí.")
    con_valor = [c for c in t['celdas'] if c['valor'] and c['valor'].get('ok') is not None]
    if len(con_valor) > 1 and not any(c['valor']['ok'] for c in con_valor):
        salida.append("En ninguna de esas semanas llegó a su objetivo según TRAX.")
    motivos = {m for m in serie['incidencia'] if m}
    if len(serie) > 1 and len(motivos) == 1 and motivo_actual:
        salida.append(f"Todas dicen lo mismo ({_limpio(motivo_actual)}). Si es el mismo problema cada semana, "
                      f"conviene resolverlo de raíz con TRAX.")
    rechazos = [m for m in serie.loc[serie['validacion_cliente'] == NO_APROBADA, 'motivo_cliente'] if m]
    if rechazos:
        salida.append(f"Antes se rechazó por: {_y(sorted(set(rechazos)))}.")
    return salida


def _limpio(motivo) -> str:
    m = texto(motivo)
    return 'Otro' if m.startswith('Otro (') else m.rstrip('.').lower()


def _y(partes) -> str:
    partes = list(partes)
    if len(partes) <= 1:
        return ''.join(partes)
    return ', '.join(partes[:-1]) + ' y ' + partes[-1]


# ------------------------------------------------------------
# Resumen y Excel
# ------------------------------------------------------------
def resumen(df: pd.DataFrame, ruta_area: dict = None) -> dict:
    """Los números de la pestaña Resumen."""
    ruta_area = ruta_area or {}
    est = df['validacion_cliente'] if len(df) else pd.Series(dtype=str)
    total = int(len(df))
    salida = {'total': total, 'pendientes': int((est == PENDIENTE).sum()),
              'aprobadas': int((est == APROBADA).sum()), 'no_aprobadas': int((est == NO_APROBADA).sum())}
    salida['validadas'] = salida['aprobadas'] + salida['no_aprobadas']

    def conteo(sub):
        e = sub['validacion_cliente']
        return {'total': int(len(sub)), 'aprobadas': int((e == APROBADA).sum()),
                'no_aprobadas': int((e == NO_APROBADA).sum()), 'pendientes': int((e == PENDIENTE).sum())}

    salida['por_kpi'] = [dict(kpi=k, **conteo(df[df['tipo'] == k])) for k in KPIS] if total else []
    areas = df['ruta'].map(lambda r: ruta_area.get(r) or 'Sin área') if total else pd.Series(dtype=str)
    salida['por_area'] = [dict(area=a, **conteo(df[areas == a])) for a in sorted(areas.unique())] if total else []
    # Lo que dijo el supervisor contra lo que decidió el cliente.
    salida['cruce'] = {(s, c): int(((df['estado'] == s) & (est == c)).sum())
                       for s in ('AUTORIZADA', 'NO_AUTORIZADA', 'PENDIENTE') for c in ESTADOS_CLIENTE} if total else {}
    motivos = df.loc[est == NO_APROBADA, 'motivo_cliente'] if total else pd.Series(dtype=str)
    salida['motivos_no'] = [(m or 'Sin motivo', int(n)) for m, n in motivos.value_counts().items()]
    ultima = pd.to_datetime(df['validada_en'], utc=True, errors='coerce').max() if total else pd.NaT
    salida['ultima'] = None if pd.isna(ultima) else ultima
    return salida


def _celda_excel(c, v):
    if c in ('fotos', 'productos'):
        return json.dumps(v, ensure_ascii=False, separators=(',', ':')) if v else None
    if c in ('id', 'semana'):
        return None if v is None or pd.isna(v) else int(v)
    if c == 'curt':
        n = numero(v)
        return int(n) if n is not None and float(n).is_integer() else (texto(v) or None)
    return texto(v) or None


def excel_finales(df: pd.DataFrame, hoja='INCIDENCIAS') -> bytes:
    """El Excel de INCIDENCIAS FINALES para el cierre, con el formato de agosto.

    Las por validar salen con VALIDACION FINAL vacía: el cuaderno las cuenta,
    avisa y las toma como NO aprobadas.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columnas = list(COLUMNAS_EXCEL) + [COLUMNA_VALIDACION] + list(COLUMNAS_EXCEL_CLIENTE)
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\[\]:*?/\\]', ' ', hoja)[:31]
    ws.append(columnas)
    filas = df.sort_values('id') if len(df) else df
    for _, f in filas.iterrows():
        fila = [_celda_excel(c, f.get(c)) for c in COLUMNAS_EXCEL]
        fila.append(TEXTO_EXCEL.get(f.get('validacion_cliente'), None))
        fila += [_celda_excel(c, f.get(c)) for c in COLUMNAS_EXCEL_CLIENTE]
        ws.append(fila)

    negrita = Font(bold=True, color='1F2A5C')
    for celda in ws[1]:
        celda.font = negrita
        celda.fill = PatternFill('solid', fgColor='EEF3FE')
        celda.alignment = Alignment(vertical='center')
    ws[1][len(COLUMNAS_EXCEL)].fill = PatternFill('solid', fgColor='FDF0F7')
    anchos = {'tienda': 34, 'comentario': 50, 'fotos': 30, 'link_trax': 30, 'incidencia': 34,
              COLUMNA_VALIDACION: 18, 'motivo_cliente': 30, 'comentario_cliente': 40}
    for i, c in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(c, 15)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
