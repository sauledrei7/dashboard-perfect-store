"""
Exhibiciones, OSA y precios del tablero del director (v25).

Solo pandas, igual que director_calc.py: recibe las tablas ya leídas (de Supabase
o de los CSV de DATAS/<MES>/output/csv_director/) y regresa diccionarios listos
para pintar y las tablas de las descargas. Las reglas (qué es una exhibición única,
qué agotado pasa a disponible, qué precio está en rango) ya vienen aplicadas desde
DATAS/generar_csv_director.py, con las mismas funciones que los reportes en Excel;
aquí solo se suma, se compara y se acomoda.

Cómo se cuenta:
- Exhibición única = una escena de TRAX; captura = cada producto reconocido en ella.
  Por categoría o marca, las únicas son las escenas con algo de esa categoría o
  marca, así que no suman el total: una escena con whisky y tequila cuenta en las dos.
- OSA = 1 - agotados finales / códigos evaluados. Agotados según TRAX = la suma de las
  10 razones; pasan a disponible las 6 de reconocimiento o acomodo.
- Adherencia = (capturas en rango ±5% + sin precio esperado) / capturas.
- Semana 0 = el mes completo. Con una semana elegida se compara contra la semana
  anterior (la última del mes anterior si es la primera); con el mes, contra el mes
  anterior.
- Las tiendas se agrupan como en el resto del tablero: el área por la ruta con la
  estructura más reciente, el supervisor por la ruta en el periodo.
"""
import io

import numpy as np
import pandas as pd

from director_calc import nombre_corto

MES = 0

# (columna, nombre, pasa a disponible). El orden es el de la tarjeta "Por qué se agota".
RAZONES = [
    ('r_sin_inventario', 'Sin inventario', False),
    ('r_no_catalogado', 'No catalogado', False),
    ('r_inventario_fantasma', 'Inventario fantasma', False),
    ('r_sin_respuesta', 'Sin respuesta del promotor', False),
    ('r_no_exhibible', 'En tienda, no puede exhibirse', True),
    ('r_nueva_imagen', 'Nueva imagen', True),
    ('r_no_reconocido', 'En anaquel, TRAX no lo reconoció', True),
    ('r_militraje', 'Militraje (ml)', True),
    ('r_en_bodega', 'En bodega, no en anaquel', True),
    ('r_obstruido', 'Obstruido o bloqueado', True),
]
COLS_RAZON = [c for c, _, _ in RAZONES]
COLS_SIGUEN = [c for c, _, pasa in RAZONES if not pasa]
COLS_PASAN = [c for c, _, pasa in RAZONES if pasa]
NOMBRE_RAZON = {c: n for c, n, _ in RAZONES}

UBICACIONES = {'ARETE': 'Arete', 'BUNKER / CAMARA FRIA': 'Búnker / cámara fría', 'CABECERA': 'Cabecera',
               'EXHIBIDORES': 'Exhibidores', 'Exhibición Temporada – Bandera': 'Temporada / bandera',
               'ISLA': 'Isla', 'PROMOCIONES': 'Promociones', 'MAMUT': 'Mamut', 'PALLET': 'Pallet'}
ORDEN_CATEGORIAS = ['Whisky', 'Tequila', 'Vodka', 'Licor', 'Ron', 'Ginebra', 'Mezcal']
MATERIAL_POS = 'Material POS'
TRAMOS_DIF = [('abajo_15', '15% o más abajo'), ('abajo_5_15', 'Entre 5% y 15% abajo'),
              ('en_rango', 'Dentro de ±5%'), ('arriba_5_15', 'Entre 5% y 15% arriba'),
              ('arriba_15', '15% o más arriba'), ('sin_esperado', 'Sin precio esperado (cuenta OK)')]
CAMPOS_PRECIO = ['capturas', 'en_rango', 'sin_esperado', 'abajo_15', 'abajo_5_15', 'arriba_5_15', 'arriba_15']
LIGA_ESCENA = 'https://services.traxretail.com/trax-one/diageomx/explore/scene/{}'
GRUPOS = {'area': 'area', 'supervisor': 'supervisor', 'ruta': 'ruta', 'cadena': 'cadena'}


# ------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------
def _num(df, campos):
    """Columnas numéricas sin nulos (Supabase manda None donde no hay dato)."""
    df = df.copy()
    for c in campos:
        df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0) if c in df.columns else 0
    return df


def _curt(df):
    df = df.copy()
    df['curt'] = df['curt'].astype(str).str.replace(r'\.0$', '', regex=True)
    return df


def _f(x):
    """Número de Python, o None si no hay dato."""
    try:
        x = float(x)
    except (TypeError, ValueError):
        return None
    return None if np.isnan(x) or np.isinf(x) else x


def _div(a, b):
    return a / b if b else None


def anterior(semanas, s, s_fin_previo):
    """Semana con la que se compara s (0 = mes): (semana, ¿es del periodo anterior?)."""
    if s == MES:
        return MES, True
    if s > semanas[0]:
        return s - 1, False
    return (s_fin_previo, True) if s_fin_previo else (None, True)


def estructura(rt: pd.DataFrame, kp: pd.DataFrame, ruta_area: dict) -> pd.DataFrame:
    """Una fila por tienda del periodo: nombre, cadena, ruta, supervisor, área y si se visitó."""
    if rt is None or len(rt) == 0:
        return pd.DataFrame(columns=['tienda', 'cadena', 'ruta', 'supervisor', 'sup', 'area', 'visitada'])
    t = _curt(rt).drop_duplicates('curt').set_index('curt')
    sup_de_ruta = kp.drop_duplicates('ruta').set_index('ruta')['supervisor'] if len(kp) else pd.Series(dtype=object)
    area_de_ruta = kp.drop_duplicates('ruta').set_index('ruta')['area_manager'] if len(kp) else pd.Series(dtype=object)
    e = pd.DataFrame({
        'tienda': t['tienda'].fillna('').astype(str),
        'cadena': t['cadena'].fillna('Sin cadena').astype(str),
        'ruta': t['ruta'].astype(str),
        'supervisor': t['ruta'].map(sup_de_ruta),
        'area': t['ruta'].map(ruta_area).fillna(t['ruta'].map(area_de_ruta)),
        'visitada': t['tienda_visitada'].eq(True) if 'tienda_visitada' in t.columns else False,
    })
    e['sup'] = e['supervisor'].map(nombre_corto)
    return e


def _etiqueta_grupo(grupo, fila):
    if grupo == 'supervisor':
        return nombre_corto(fila['k']), fila.get('area') or ''
    if grupo == 'ruta':
        return fila['k'], f"{nombre_corto(fila.get('supervisor'))} · {fila.get('area') or ''}"
    return fila['k'], ''


# ------------------------------------------------------------
# Exhibiciones
# ------------------------------------------------------------
CAMPOS_EXH = ['exhibiciones', 'capturas', 'frentes', 'alto_impacto', 'sin_respuesta']


def _corte_mes(corte, nombre):
    c = _num(corte, ['exhibiciones', 'capturas', 'frentes', 'tiendas', 'alto_impacto', 'semana'])
    return c[(c['semana'] == MES) & (c['corte'] == nombre)]


def _con_previo(actual, previo, llave='clave'):
    p = previo.set_index(llave) if len(previo) else pd.DataFrame()
    filas = []
    for _, x in actual.iterrows():
        a = p.loc[x[llave]] if len(p) and x[llave] in p.index else None
        filas.append({'k': x[llave], 'padre': x.get('padre') or '', 'exh': int(x['exhibiciones']), 'cap': int(x['capturas']),
                      'fre': int(x['frentes']), 'tiendas': int(x['tiendas']), 'alto': int(x['alto_impacto']),
                      'exh_ant': None if a is None else int(a['exhibiciones']),
                      'cap_ant': None if a is None else int(a['capturas'])})
    return filas


def exh_vista(ts, corte, ts_prev, corte_prev, est, semanas) -> dict:
    """Todo lo que pinta la sección Exhibiciones de un mes."""
    if corte is None or len(corte) == 0:
        return None
    total = _corte_mes(corte, 'total')
    if len(total) == 0:
        return None
    T = total.iloc[0]
    Tp = _corte_mes(corte_prev, 'total') if corte_prev is not None and len(corte_prev) else pd.DataFrame()
    Tp = Tp.iloc[0] if len(Tp) else None
    n_tiendas = len(est)
    tot = {
        'exh': int(T['exhibiciones']), 'cap': int(T['capturas']), 'fre': int(T['frentes']),
        'alto': int(T['alto_impacto']), 'tiendas_con': int(T['tiendas']), 'maestro': n_tiendas,
        'exh_ant': None if Tp is None else int(Tp['exhibiciones']),
        'cap_ant': None if Tp is None else int(Tp['capturas']),
        'fre_ant': None if Tp is None else int(Tp['frentes']),
        'alto_pct_ant': None if Tp is None else _div(Tp['alto_impacto'], Tp['exhibiciones']),
    }
    tot['alto_pct'] = _div(tot['alto'], tot['exh'])
    tot['por_tienda'] = _div(tot['exh'], n_tiendas)
    tot['cap_por_tienda'] = _div(tot['cap'], n_tiendas)
    imp = _corte_mes(corte, 'impacto').set_index('clave')['exhibiciones'] if len(corte) else pd.Series(dtype=float)
    tot['sin_respuesta'] = int(imp.get('Sin respuesta', 0))

    # Un mes anterior sin datos del tablero (junio, o uno al que le faltan sus CSV) llega de
    # Supabase como tabla vacía y sin columnas: se toma como si no hubiera mes anterior.
    prev = corte_prev if corte_prev is not None and len(corte_prev) else pd.DataFrame(columns=corte.columns)
    cats = _corte_mes(corte, 'categoria')
    pos = cats[cats['clave'] == MATERIAL_POS]
    cats = cats[cats['clave'] != MATERIAL_POS].copy()
    cats['_o'] = cats['clave'].map({c: i for i, c in enumerate(ORDEN_CATEGORIAS)}).fillna(99)
    categorias = _con_previo(cats.sort_values(['_o', 'capturas'], ascending=[True, False]), _corte_mes(prev, 'categoria'))

    marcas = _con_previo(_corte_mes(corte, 'marca').sort_values('capturas', ascending=False), _corte_mes(prev, 'marca'))
    variantes = _con_previo(_corte_mes(corte, 'variante').sort_values('capturas', ascending=False), _corte_mes(prev, 'variante'))
    por_marca = {}
    for v in variantes:
        por_marca.setdefault(v['padre'], []).append(v)
    for m in marcas:
        m['variantes'] = por_marca.get(m['k'], [])
        m['cat'] = m.pop('padre')

    ubic = _con_previo(_corte_mes(corte, 'ubicacion').sort_values('exhibiciones', ascending=False), _corte_mes(prev, 'ubicacion'))
    for u in ubic:
        u['nombre'] = UBICACIONES.get(u['k'], u['k'])

    c = _num(corte, ['exhibiciones', 'capturas', 'alto_impacto', 'semana'])
    sem = c[(c['corte'] == 'total') & (c['semana'] != MES)].set_index('semana')
    semanas_v = [{'s': s, 'exh': int(sem.loc[s, 'exhibiciones']) if s in sem.index else 0,
                  'cap': int(sem.loc[s, 'capturas']) if s in sem.index else 0,
                  'alto_pct': _div(sem.loc[s, 'alto_impacto'], sem.loc[s, 'exhibiciones']) if s in sem.index else None}
                 for s in semanas]

    # Por tienda: el mes es la suma de las semanas
    t = _curt(_num(ts, CAMPOS_EXH)).groupby('curt')[CAMPOS_EXH].sum() if len(ts) else pd.DataFrame(columns=CAMPOS_EXH)
    tiendas = est.join(t, how='left').fillna({k: 0 for k in CAMPOS_EXH})
    tp = _curt(_num(ts_prev, CAMPOS_EXH)).groupby('curt')[CAMPOS_EXH].sum() if ts_prev is not None and len(ts_prev) else None

    cortes_dist = [(0, 0, '0'), (1, 4, '1–4'), (5, 8, '5–8'), (9, 12, '9–12'), (13, 16, '13–16'), (17, 24, '17–24'),
                   (25, 10 ** 9, '25 o más')]
    dist = [{'k': k, 'n': int(tiendas['exhibiciones'].between(a, b).sum())} for a, b, k in cortes_dist]

    def equipo(grupo):
        g = tiendas.groupby(GRUPOS[grupo])
        tabla = pd.DataFrame({'tiendas': g.size(), 'con': g['exhibiciones'].apply(lambda s: int((s > 0).sum())),
                              'exh': g['exhibiciones'].sum(), 'cap': g['capturas'].sum(), 'fre': g['frentes'].sum(),
                              'alto': g['alto_impacto'].sum(), 'area': g['area'].first(), 'supervisor': g['supervisor'].first()})
        previo = None
        if tp is not None:
            ant = est.join(tp, how='inner')
            previo = ant.groupby(GRUPOS[grupo]).agg(exh=('exhibiciones', 'sum'), n=('exhibiciones', 'size'))
        filas = []
        for k, x in tabla.iterrows():
            fila = {'k': k, 'tiendas': int(x['tiendas']), 'sin': int(x['tiendas'] - x['con']), 'exh': int(x['exh']),
                    'cap': int(x['cap']), 'fre': int(x['fre']), 'alto_pct': _div(x['alto'], x['exh']),
                    'por': _div(x['exh'], x['tiendas']), 'area': x['area'], 'supervisor': x['supervisor']}
            if previo is not None and k in previo.index:
                fila['por_ant'] = _div(previo.loc[k, 'exh'], x['tiendas'])
            fila['nombre'], fila['sub'] = _etiqueta_grupo(grupo, fila)
            filas.append(fila)
        return filas

    sin = tiendas[tiendas['exhibiciones'] == 0]
    sin_cadena = (sin.groupby('cadena').agg(n=('exhibiciones', 'size'), vis=('visitada', 'sum'))
                  .sort_values(['n', 'vis'], ascending=False).reset_index())
    mas = tiendas[tiendas['exhibiciones'] > 0].sort_values(['exhibiciones', 'capturas'], ascending=False).head(8)
    return {
        'tot': tot,
        'categorias': categorias, 'suma_categorias': int(sum(c['exh'] for c in categorias)),
        'pos': int(pos['capturas'].sum()) if len(pos) else 0,
        'marcas': marcas, 'variantes': variantes, 'ubicaciones': ubic, 'semanas': semanas_v, 'dist': dist,
        'equipo': {g: equipo(g) for g in ('area', 'supervisor', 'cadena')},
        'sin': {'n': len(sin), 'vis': int(sin['visitada'].sum()),
                'por_cadena': sin_cadena.head(8).to_dict('records'),
                'visitadas': [{'curt': k, **x} for k, x in sin[sin['visitada']].sort_values(['cadena', 'tienda']).head(10)
                              [['tienda', 'cadena', 'sup', 'ruta']].iterrows()]},
        'mas': [{'curt': k, 'tienda': x['tienda'], 'cadena': x['cadena'], 'sup': x['sup'], 'ruta': x['ruta'],
                 'exh': int(x['exhibiciones']), 'cap': int(x['capturas'])} for k, x in mas.iterrows()],
    }


def exh_detalle_tienda(escenas: pd.DataFrame) -> list:
    """Las exhibiciones de una tienda, de la más vieja a la más nueva, con su liga a TRAX."""
    if escenas is None or len(escenas) == 0:
        return []
    e = _num(escenas, ['capturas', 'frentes', 'semana'])
    e = e.sort_values(['fecha', 'ubicacion', 'scene_id'])
    return [{'fecha': pd.Timestamp(x['fecha']).strftime('%d/%m'), 'semana': int(x['semana']),
             'ubicacion': UBICACIONES.get(x['ubicacion'], x['ubicacion']), 'impacto': x['impacto'],
             'categoria': x['categoria'], 'variante': x['variante'], 'productos': x['productos'],
             'capturas': int(x['capturas']), 'frentes': int(x['frentes']),
             'liga': LIGA_ESCENA.format(int(x['scene_id']))} for _, x in e.iterrows()]


def tabla_exh_tiendas(ts, est, semanas) -> pd.DataFrame:
    """Descarga: cada tienda con sus exhibiciones únicas y capturas de cada semana y del mes."""
    t = _curt(_num(ts, CAMPOS_EXH + ['semana']))
    base = est[['tienda', 'cadena', 'area', 'sup', 'ruta']].copy()
    for s in semanas:
        x = t[t['semana'] == s].set_index('curt')
        base[f'Exhibiciones S{s}'] = x['exhibiciones'].reindex(base.index).fillna(0).astype(int)
        base[f'Capturas S{s}'] = x['capturas'].reindex(base.index).fillna(0).astype(int)
    mes = t.groupby('curt')[CAMPOS_EXH].sum().reindex(base.index).fillna(0).astype(int)
    base['Exhibiciones del mes'] = mes['exhibiciones']
    base['Capturas del mes'] = mes['capturas']
    base['Frentes del mes'] = mes['frentes']
    base['Alto impacto'] = mes['alto_impacto']
    base = base.rename(columns={'tienda': 'Tienda', 'cadena': 'Cadena', 'area': 'Área', 'sup': 'Supervisor', 'ruta': 'Promotor'})
    return base.rename_axis('CURT').reset_index().sort_values(['Exhibiciones del mes', 'CURT'])


# ------------------------------------------------------------
# OSA
# ------------------------------------------------------------
def _osa(df):
    df = _num(df, ['codigos', 'semana'] + COLS_RAZON)
    df['at'] = df[COLS_RAZON].sum(axis=1)
    df['rc'] = df[COLS_PASAN].sum(axis=1)
    df['ag'] = df[COLS_SIGUEN].sum(axis=1)
    return df


def _del_periodo(df, s, llave):
    """Filas de una semana, o del mes completo (la suma de las semanas)."""
    if df is None or len(df) == 0:
        return None
    x = df if s == MES else df[df['semana'] == s]
    if len(x) == 0:
        return None
    campos = [c for c in x.columns if c not in (llave, 'semana', 'periodo_id', 'id', 'created_at', 'marca', 'categoria')
              and pd.api.types.is_numeric_dtype(x[c])]
    extra = [c for c in ('marca', 'categoria') if c in x.columns]
    g = x.groupby(llave)
    out = g[campos].sum()
    for c in extra:
        out[c] = g[c].first()
    return out


def _razon_principal(fila):
    valores = [(fila.get(c, 0) or 0, c) for c in COLS_SIGUEN]
    n, c = max(valores)
    return NOMBRE_RAZON[c] if n > 0 else ''


def _osa_valor(o):
    return None if o is None or not o.get('codigos') else 1 - o['ag'] / o['codigos']


def _suma(df, campos):
    return {c: float(df[c].sum()) for c in campos} if df is not None and len(df) else None


def osa_vista(ots, osku, ots_prev, osku_prev, est, semanas, s, s_fin_previo) -> dict:
    """La sección OSA para la semana s (0 = mes completo)."""
    if ots is None or len(ots) == 0:
        return None
    ots = _curt(_osa(ots))
    osku = _osa(osku) if osku is not None and len(osku) else None
    ots_prev = _curt(_osa(ots_prev)) if ots_prev is not None and len(ots_prev) else None
    osku_prev = _osa(osku_prev) if osku_prev is not None and len(osku_prev) else None
    s_ant, del_previo = anterior(semanas, s, s_fin_previo)
    base_t = ots_prev if del_previo else ots
    base_p = osku_prev if del_previo else osku
    campos = ['codigos', 'at', 'rc', 'ag'] + COLS_RAZON

    # Solo tiendas de la estructura del periodo, igual que el resto del tablero
    act = _del_periodo(ots, s, 'curt')
    if act is None:
        return None
    act = act[act.index.isin(est.index)]
    ant = _del_periodo(base_t, s_ant, 'curt') if base_t is not None and s_ant is not None else None
    if ant is not None:
        ant = ant[ant.index.isin(est.index)]
    cols_est = ['area', 'supervisor', 'ruta', 'cadena']
    tot = _suma(act, campos)
    tot['tiendas'] = len(act)
    tot_ant = _suma(ant, campos)

    sem = []
    for w in semanas:
        x = _del_periodo(ots, w, 'curt')
        sem.append({'s': w, 'osa': _osa_valor(_suma(x, campos)) if x is not None else None})

    razones = [{'k': n, 'pasa': pasa, 'n': int(tot[c]), 'pct': _div(tot[c], tot['codigos']),
                'pct_ant': _div(tot_ant[c], tot_ant['codigos']) if tot_ant else None}
               for c, n, pasa in RAZONES if tot[c] > 0]

    def productos(nivel):
        if osku is None:
            return []
        llave = {'sku': 'sku', 'marca': 'marca', 'categoria': 'categoria'}[nivel]
        x = _del_periodo(osku, s, llave)
        if x is None:
            return []
        y = _del_periodo(base_p, s_ant, llave) if base_p is not None and s_ant is not None else None
        filas = []
        for k, f in x.iterrows():
            d = f.to_dict()
            a = y.loc[k].to_dict() if y is not None and k in y.index else None
            filas.append({'k': k, 'marca': d.get('marca', ''), 'categoria': d.get('categoria', ''),
                          'codigos': int(d['codigos']), 'at': int(d['at']), 'rc': int(d['rc']), 'ag': int(d['ag']),
                          'osa': _osa_valor(d), 'osa_ant': _osa_valor(a), 'razon': _razon_principal(d),
                          'skus': int((osku[osku[llave] == k]['sku'].nunique()) if llave != 'sku' else 1)})
        return sorted(filas, key=lambda r: (r['osa'] if r['osa'] is not None else 2))

    def equipo(grupo):
        col = GRUPOS[grupo]
        g = act.join(est[cols_est], how='inner').groupby(col)
        tabla = g[campos].sum()
        tabla['tiendas'] = g.size()
        tabla['area'] = g['area'].first()
        tabla['supervisor'] = g['supervisor'].first()
        previo = None
        if ant is not None:
            previo = ant.join(est[cols_est], how='inner').groupby(col)[campos].sum()
        filas = []
        for k, f in tabla.iterrows():
            d = f.to_dict()
            fila = {'k': k, 'tiendas': int(d['tiendas']), 'codigos': int(d['codigos']), 'ag': int(d['ag']),
                    'osa': _osa_valor(d), 'razon': _razon_principal(d), 'area': d['area'], 'supervisor': d['supervisor']}
            if previo is not None and k in previo.index:
                fila['osa_ant'] = _osa_valor(previo.loc[k].to_dict())
            fila['nombre'], fila['sub'] = _etiqueta_grupo(grupo, fila)
            filas.append(fila)
        return sorted(filas, key=lambda r: (r['osa'] if r['osa'] is not None else 2))

    tiendas = act.join(est[['tienda', 'cadena', 'sup', 'ruta']], how='inner')
    filas_t = []
    for k, f in tiendas.iterrows():
        d = f.to_dict()
        a = ant.loc[k].to_dict() if ant is not None and k in ant.index else None
        filas_t.append({'curt': k, 'tienda': d['tienda'], 'cadena': d['cadena'], 'sup': d['sup'], 'ruta': d['ruta'],
                        'codigos': int(d['codigos']), 'ag': int(d['ag']), 'osa': _osa_valor(d), 'osa_ant': _osa_valor(a),
                        'razon': _razon_principal(d)})
    filas_t.sort(key=lambda r: (r['osa'] if r['osa'] is not None else 2, -r['codigos']))
    return {
        'tot': {**tot, 'osa': _osa_valor(tot), 'osa_trax': 1 - tot['at'] / tot['codigos'] if tot['codigos'] else None,
                'osa_ant': _osa_valor(tot_ant), 'razon': _razon_principal(tot)},
        's_ant': s_ant, 'del_previo': del_previo,
        'semanas': sem, 'razones': razones,
        'productos': {n: productos(n) for n in ('categoria', 'marca', 'sku')},
        'equipo': {g: equipo(g) for g in ('area', 'supervisor', 'ruta', 'cadena')},
        'tiendas': filas_t,
    }


def tabla_osa_tiendas(ots, est, s) -> pd.DataFrame:
    """Descarga: cada tienda en la semana s (0 = el mes), del OSA más bajo al más alto."""
    x = _del_periodo(_curt(_osa(ots)), s, 'curt')
    if x is None:
        return pd.DataFrame()
    x = x.join(est[['tienda', 'cadena', 'area', 'sup', 'ruta']], how='inner')
    out = pd.DataFrame({
        'CURT': x.index, 'Tienda': x['tienda'], 'Cadena': x['cadena'], 'Área': x['area'], 'Supervisor': x['sup'],
        'Promotor': x['ruta'], 'Periodo': 'Mes completo' if s == MES else f'S{s}',
        'Códigos evaluados': x['codigos'].astype(int), 'Agotados según TRAX': x['at'].astype(int),
        'Pasan a disponible': x['rc'].astype(int), 'Siguen agotados': x['ag'].astype(int),
        '% OSA': 1 - x['ag'] / x['codigos'], '% OSA según TRAX': 1 - x['at'] / x['codigos'],
        'Razón principal': [_razon_principal(f.to_dict()) for _, f in x.iterrows()],
    })
    for c, n, _ in RAZONES:
        out[n] = x[c].astype(int)
    return out.sort_values(['% OSA', 'Códigos evaluados'], ascending=[True, False])


def tabla_osa_final(ots, est, semanas) -> pd.DataFrame:
    """Descarga: resultado final del mes por tienda, con el OSA de cada semana."""
    base = tabla_osa_tiendas(ots, est, MES)
    if len(base) == 0:
        return base
    o = _curt(_osa(ots))
    for s in semanas:
        x = o[o['semana'] == s].groupby('curt')[['codigos', 'ag']].sum()
        base.insert(6 + semanas.index(s), f'OSA S{s}', (1 - x['ag'] / x['codigos']).reindex(base['CURT']).to_numpy())
    return base.drop(columns='Periodo').rename(columns={'% OSA': '% OSA del mes', 'Códigos evaluados': 'Códigos del mes'})


# ------------------------------------------------------------
# Precios
# ------------------------------------------------------------
def _pre(df):
    df = _num(df, CAMPOS_PRECIO + ['semana'])
    df['arriba'] = df['arriba_5_15'] + df['arriba_15']
    df['abajo'] = df['abajo_15'] + df['abajo_5_15']
    df['ok'] = df['en_rango'] + df['sin_esperado']
    return df


def _adh(o):
    return None if o is None or not o.get('capturas') else o['ok'] / o['capturas']


def _mediana(corte, s, nombre, clave):
    if corte is None or len(corte) == 0:
        return None
    x = corte[(pd.to_numeric(corte['semana'], errors='coerce') == s) & (corte['corte'] == nombre) & (corte['clave'] == str(clave))]
    return _f(x['dif_mediana'].iloc[0]) if len(x) else None


def precios_vista(pts, pcorte, pts_prev, pcorte_prev, est, semanas, s, s_fin_previo) -> dict:
    """La sección Precios para la semana s (0 = mes completo)."""
    if pts is None or len(pts) == 0:
        return None
    pts = _curt(_pre(pts))
    pts_prev = _curt(_pre(pts_prev)) if pts_prev is not None and len(pts_prev) else None
    s_ant, del_previo = anterior(semanas, s, s_fin_previo)
    base_t = pts_prev if del_previo else pts
    base_c = pcorte_prev if del_previo else pcorte
    campos = CAMPOS_PRECIO + ['arriba', 'abajo', 'ok']

    act = _del_periodo(pts, s, 'curt')
    if act is None:
        return None
    act = act.join(est[['tienda', 'cadena', 'ruta', 'supervisor', 'sup', 'area']], how='inner')
    ant = _del_periodo(base_t, s_ant, 'curt') if base_t is not None and s_ant is not None else None
    if ant is not None:
        ant = ant.join(est[['cadena', 'ruta', 'supervisor', 'area']], how='inner')
    tot = _suma(act, campos)
    tot['tiendas'] = len(act)
    tot_ant = _suma(ant, campos)

    sem = []
    for w in semanas:
        x = _del_periodo(pts, w, 'curt')
        sem.append({'s': w, 'adh': _adh(_suma(x, campos)) if x is not None else None})

    def productos(nivel):
        if pcorte is None or len(pcorte) == 0:
            return []
        c = _num(pcorte, ['semana', 'capturas', 'en_rango', 'sin_esperado', 'arriba', 'abajo'])
        x = c[(c['semana'] == s) & (c['corte'] == nivel)]
        y = None
        if base_c is not None and len(base_c) and s_ant is not None:
            cb = _num(base_c, ['semana', 'capturas', 'en_rango', 'sin_esperado', 'arriba', 'abajo'])
            y = cb[(cb['semana'] == s_ant) & (cb['corte'] == nivel)].set_index('clave')
        filas = []
        for _, f in x.iterrows():
            d = f.to_dict()
            d['ok'] = d['en_rango'] + d['sin_esperado']
            a = None
            if y is not None and d['clave'] in y.index:
                a = y.loc[d['clave']].to_dict()
                a['ok'] = a['en_rango'] + a['sin_esperado']
            skus = None
            if nivel != 'sku':
                todos = c[(c['semana'] == s) & (c['corte'] == 'sku')]
                skus = int((todos[nivel] == d['clave']).sum())
            filas.append({'k': d['clave'], 'marca': d.get('marca') or '', 'categoria': d.get('categoria') or '', 'skus': skus,
                          'capturas': int(d['capturas']), 'adh': _adh(d), 'adh_ant': _adh(a),
                          'arriba': _div(d['arriba'], d['capturas']), 'abajo': _div(d['abajo'], d['capturas']),
                          'precio_estrategia': _f(d.get('precio_estrategia')), 'precio_tipico': _f(d.get('precio_tipico')),
                          'dif': _f(d.get('dif_mediana'))})
        return sorted(filas, key=lambda r: (r['adh'] if r['adh'] is not None else 2))

    def equipo(grupo):
        col = GRUPOS[grupo]
        g = act.groupby(col)
        tabla = g[campos].sum()
        tabla['tiendas'] = g.size()
        tabla['area'] = g['area'].first()
        tabla['supervisor'] = g['supervisor'].first()
        previo = ant.groupby(col)[campos].sum() if ant is not None else None
        filas = []
        for k, f in tabla.iterrows():
            d = f.to_dict()
            fila = {'k': k, 'tiendas': int(d['tiendas']), 'capturas': int(d['capturas']), 'adh': _adh(d),
                    'arriba': _div(d['arriba'], d['capturas']), 'abajo': _div(d['abajo'], d['capturas']),
                    'dif': _mediana(pcorte, s, grupo, k), 'area': d['area'], 'supervisor': d['supervisor']}
            if previo is not None and k in previo.index:
                fila['adh_ant'] = _adh(previo.loc[k].to_dict())
            fila['nombre'], fila['sub'] = _etiqueta_grupo(grupo, fila)
            filas.append(fila)
        return sorted(filas, key=lambda r: (r['adh'] if r['adh'] is not None else 2))

    filas_t = []
    for k, f in act.iterrows():
        d = f.to_dict()
        a = ant.loc[k].to_dict() if ant is not None and k in ant.index else None
        filas_t.append({'curt': k, 'tienda': d['tienda'], 'cadena': d['cadena'], 'sup': d['sup'], 'ruta': d['ruta'],
                        'capturas': int(d['capturas']), 'adh': _adh(d), 'adh_ant': _adh(a),
                        'arriba': _div(d['arriba'], d['capturas']), 'abajo': _div(d['abajo'], d['capturas'])})
    filas_t.sort(key=lambda r: (r['adh'] if r['adh'] is not None else 2, -r['capturas']))
    return {
        'tot': {**tot, 'adh': _adh(tot), 'adh_ant': _adh(tot_ant), 'arriba_pct': _div(tot['arriba'], tot['capturas']),
                'abajo_pct': _div(tot['abajo'], tot['capturas']), 'dif': _mediana(pcorte, s, 'total', 'total')},
        's_ant': s_ant, 'del_previo': del_previo,
        'semanas': sem,
        'tramos': [{'k': n, 'campo': c, 'n': int(tot[c]), 'pct': _div(tot[c], tot['capturas'])} for c, n in TRAMOS_DIF],
        'productos': {n: productos(n) for n in ('categoria', 'marca', 'sku')},
        'equipo': {g: equipo(g) for g in ('cadena', 'area', 'supervisor', 'ruta')},
        'tiendas': filas_t,
    }


def tabla_precios_tiendas(pts, est, s) -> pd.DataFrame:
    """Descarga: cada tienda en la semana s (0 = el mes), de la menor adherencia a la mayor."""
    x = _del_periodo(_curt(_pre(pts)), s, 'curt')
    if x is None:
        return pd.DataFrame()
    x = x.join(est[['tienda', 'cadena', 'area', 'sup', 'ruta']], how='inner')
    out = pd.DataFrame({
        'CURT': x.index, 'Tienda': x['tienda'], 'Cadena': x['cadena'], 'Área': x['area'], 'Supervisor': x['sup'],
        'Promotor': x['ruta'], 'Periodo': 'Mes completo' if s == MES else f'S{s}',
        'Capturas': x['capturas'].astype(int), 'En rango ±5%': x['en_rango'].astype(int),
        'Sin precio esperado (OK)': x['sin_esperado'].astype(int), 'Arriba +5%': x['arriba'].astype(int),
        'Abajo -5%': x['abajo'].astype(int), '% adherencia': x['ok'] / x['capturas'],
        '15% o más arriba': x['arriba_15'].astype(int), '15% o más abajo': x['abajo_15'].astype(int),
    })
    return out.sort_values(['% adherencia', 'Capturas'], ascending=[True, False])


def tabla_precios_final(pts, est, semanas) -> pd.DataFrame:
    """Descarga: resultado final del mes por tienda, con la adherencia de cada semana."""
    base = tabla_precios_tiendas(pts, est, MES)
    if len(base) == 0:
        return base
    p = _curt(_pre(pts))
    for s in semanas:
        x = p[p['semana'] == s].groupby('curt')[['ok', 'capturas']].sum()
        base.insert(6 + semanas.index(s), f'Adherencia S{s}', (x['ok'] / x['capturas']).reindex(base['CURT']).to_numpy())
    return base.drop(columns='Periodo').rename(columns={'% adherencia': '% adherencia del mes', 'Capturas': 'Capturas del mes'})


# ------------------------------------------------------------
# Excel de las descargas
# ------------------------------------------------------------
def excel(df: pd.DataFrame, hoja: str, titulo: str) -> bytes:
    """Una hoja con renglón de título, encabezado de color, filtros y % con formato."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = hoja[:31]
    ws.append([titulo])
    ws['A1'].font = Font(bold=True)
    ws.append(list(df.columns))
    for j in range(1, len(df.columns) + 1):
        c = ws.cell(row=2, column=j)
        c.fill = PatternFill('solid', fgColor='1F2A5C')
        c.font = Font(bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for fila in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, 'item') else v) for v in fila])
    for j, col in enumerate(df.columns, start=1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = 38 if col == 'Tienda' else (14 if len(str(col)) < 14 else 18)
        if str(col).startswith(('%', 'OSA S', 'Adherencia S')):
            for i in range(3, len(df) + 3):
                ws.cell(row=i, column=j).number_format = '0.0%'
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = 'C3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(df.columns))}{len(df) + 2}'
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
